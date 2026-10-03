from io import BytesIO
from pathlib import Path
import json

from openpyxl import load_workbook
import pytest

from app import app
from summary_corrections import (
    corrected_effect_heterogeneity_plot,
    corrected_rows,
    summary_mode,
)


ROOT = Path(__file__).resolve().parents[1]


def _reviewed_processed_headline():
    from automation.sync_sources import assembled_source
    path = assembled_source(ROOT, 'metaPee')/'combined/provisional_ui/subexposure_views/no/Cached_results/processed_meat_intake/bladder_cancer_incidence_true_all.json'
    return json.loads(path.read_text())['headline']


def _row(rows, exposure):
    return next(row for row in rows if row['Exposure'] == exposure)


def test_parent_scope_diagnostics_come_from_complete_saved_headline():
    _book, rows, disease, corrected, _source = corrected_rows(
        ROOT / 'sources/metaPee', 'no', 'bladder_cancer')
    assert disease == 'Bladder cancer'
    assert corrected == 55 + 1  # Declared reviewed UPF addition.
    processed = _row(rows, 'processed_meat_intake')
    assert processed['number estimates'] == 7 + 1  # Additional reviewed source contrast.
    saved = _reviewed_processed_headline()
    assert processed['I^2 (%)'] == pytest.approx(saved['i2'])
    assert processed['eggers p-value'] == pytest.approx(saved['eggers_p'])
    assert _row(rows, 'lycopene')['I^2 (%)'] == 0
    assert _row(rows, 'zeaxanthin')['I^2 (%)'] is None


@pytest.mark.parametrize('value,expected', [
    (None, 'yes'), ('yes', 'yes'), ('true', 'yes'), ('1', 'yes'),
    ('no', 'no'), ('false', 'no'), ('0', 'no'),
])
def test_summary_mode_matches_public_query_values(value, expected):
    assert summary_mode(value) == expected


def test_corrected_workbook_and_effect_plot_are_served():
    client = app.test_client()
    workbook = client.get(
        '/evidence/urinary/summary/workbooks/bladder_cancer/combined?include_subexposures=no')
    assert workbook.status_code == 200
    book = load_workbook(BytesIO(workbook.data), read_only=True, data_only=True)
    values = list(book.active.iter_rows(values_only=True))
    header = values[0]
    processed = dict(zip(header, next(row for row in values[1:] if row[0] == 'processed_meat_intake')))
    saved = _reviewed_processed_headline()
    assert processed['I^2 (%)'] == pytest.approx(saved['i2'])
    assert processed['eggers p-value'] == pytest.approx(saved['eggers_p'])

    plot = client.get(
        '/evidence/urinary/summary/plots/bladder_cancer/effect-heterogeneity?include_subexposures=no')
    assert plot.status_code == 200
    assert plot.headers['Content-Type'].startswith('application/pdf')
    assert plot.data.startswith(b'%PDF-')

    localized = corrected_effect_heterogeneity_plot(
        ROOT, 'no', 'bladder_cancer', 'nl')
    companion = localized.with_name('.' + localized.stem + '.tmp.png')
    assert not companion.exists()


def test_invalid_summary_mode_and_scope_are_rejected():
    client = app.test_client()
    assert client.get(
        '/evidence/urinary/summary/workbooks/bladder_cancer/combined?include_subexposures=maybe'
    ).status_code == 400
    assert client.get(
        '/evidence/urinary/summary/workbooks/not_a_scope/combined'
    ).status_code == 404
