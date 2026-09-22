"""Audit and remove superseded summary artifacts after canonical validation."""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from summary_outputs import declared_artifacts, is_summary_artifact, sha

REVISION = 'summary-consolidation-2026-09-22'
AUDIT = ROOT / 'evidence-revisions' / REVISION


def metadata(path):
    return {'path': str(path.relative_to(ROOT)), 'sha256': sha(path), 'bytes': path.stat().st_size}


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + '\n')


def old_artifacts():
    paths = subprocess.check_output(['rg', '--files', '-uuu'], cwd=ROOT, text=True).splitlines()
    selected = []
    for name in paths:
        path = ROOT / name
        if Path(name).parts[0] not in {'sources', 'runtime', 'outputs', 'updates'}:
            continue
        if is_summary_artifact(path) or (
            path.name == 'summary_manifest.json' and 'Plot' in path.parts):
            selected.append(path)
    return sorted(selected)


def prepare():
    if (AUDIT / 'inventory.json').exists():
        raise ValueError('The migration inventory already exists; do not replace its provenance')
    AUDIT.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(ROOT / 'evidence-manifest.json', AUDIT / 'prior-release-manifest.json')
    files = [metadata(path) for path in old_artifacts()]
    original_workbooks = []
    from openpyxl import load_workbook
    for source, pattern in (
        ('metaPee', 'combined/provisional_ui/subexposure_views/{mode}/Plot'),
        ('MetaVir', 'Plot/subexposures/{mode}'),
    ):
        for mode in ('yes', 'no'):
            folder = ROOT / 'sources' / source / pattern.format(mode=mode)
            manifest = json.loads((folder / 'summary_manifest.json').read_text())
            for scope in manifest['scopes'].values():
                for workbook in scope['workbooks'].values():
                    path = folder / workbook['filename']
                    book = load_workbook(path, read_only=True, data_only=False)
                    original_workbooks.append({**metadata(path), 'mode': mode,
                        'target': f'summary/{mode}/{path.name}',
                        'rows': list(book.active.values)})
                    book.close()
    write(AUDIT / 'inventory.json', {'files': files, 'workbooks': original_workbooks})
    print(f'Inventoried {len(files)} old summary artifacts and {len(original_workbooks)} current workbooks')


def validate():
    from openpyxl import load_workbook
    from pypdf import PdfReader
    reports = {}
    for mode in ('yes', 'no'):
        folder = ROOT / 'summary' / mode
        manifest = json.loads((folder / 'summary_manifest.json').read_text())
        if set(manifest.get('sources', {})) != {'metaPee', 'MetaVir'}:
            raise ValueError(f'Missing disease collection in {mode}')
        expected = declared_artifacts(manifest)
        actual = {str(path.relative_to(folder)) for path in folder.rglob('*')
                  if path.is_file() and path.suffix in {'.xlsx', '.pdf'}}
        if actual != expected:
            raise ValueError(f'Undeclared/missing artifacts: {actual ^ expected}')
        for relative in sorted(expected):
            path = folder / relative
            if sha(path) != manifest['artifacts'][relative]['sha256']:
                raise ValueError(f'Invalid checksum: {path}')
            if path.suffix == '.xlsx':
                with zipfile.ZipFile(path) as archive:
                    if archive.testzip() is not None:
                        raise ValueError(f'Invalid workbook: {path}')
                book = load_workbook(path, read_only=True, data_only=False)
                if book.active.max_column != 11:
                    raise ValueError(f'Unexpected summary columns: {path}')
                errors = [cell.coordinate for row in book.active for cell in row if cell.data_type == 'e']
                if errors:
                    raise ValueError(f'Excel errors: {path}: {errors}')
                book.close()
            else:
                pdf = PdfReader(path)
                if not pdf.pages or any(float(page.mediabox.width) <= 0 for page in pdf.pages):
                    raise ValueError(f'Invalid PDF: {path}')
        reports[mode] = {'scopes': len(manifest['scopes']),
                         'xlsx': sum(name.endswith('.xlsx') for name in expected),
                         'pdf': sum(name.endswith('.pdf') for name in expected)}
    differences = []
    inventory = json.loads((AUDIT / 'inventory.json').read_text())
    for original in inventory['workbooks']:
        target = ROOT / original['target']
        book = load_workbook(target, read_only=True, data_only=False)
        after = [list(row) for row in book.active.values]
        book.close()
        before = original['rows']
        if before[0] != after[0] or len(before) != len(after):
            raise ValueError(f'Summary structure changed: {target}')
        old_rows = {row[0]: row for row in before[1:]}
        new_rows = {row[0]: row for row in after[1:]}
        changes = []
        for exposure in sorted(set(old_rows) | set(new_rows)):
            old, new = old_rows.get(exposure), new_rows.get(exposure)
            if old != new:
                cells = [before[0][index] for index in range(len(old)) if old[index] != new[index]] if old and new else ['row membership']
                if set(cells) - {'I^2 (%)', 'eggers p-value'}:
                    raise ValueError(f'Unexpected scientific change: {target}: {exposure}: {cells}')
                changes.append({'exposure': exposure, 'columns': cells, 'before': old, 'after': new})
        if changes:
            differences.append({'path': original['target'], 'changes': changes})
    result = {'modes': reports, 'workbook_differences': differences}
    write(AUDIT / 'validation.json', result)
    print(json.dumps({'modes': reports, 'changed_workbooks': len(differences)}))
    return result


def finish():
    validation = validate()
    inventory = json.loads((AUDIT / 'inventory.json').read_text())
    expected = {item['path']: item for item in inventory['files']}
    current = {str(path.relative_to(ROOT)): path for path in old_artifacts()}
    if set(current) - set(expected):
        raise ValueError('New legacy summary artifacts appeared after inventory')
    for relative, path in current.items():
        if sha(path) != expected[relative]['sha256']:
            raise ValueError(f'An old artifact changed during the migration: {relative}')
    prior = json.loads((AUDIT / 'prior-release-manifest.json').read_text())
    manifest = json.loads((ROOT / 'evidence-manifest.json').read_text())
    changed = []
    kept = []
    for entry in manifest['files']:
        path = ROOT / entry['path']
        if entry['path'] in expected:
            changed.append({'path': entry['path'], 'before': entry, 'after': None})
            continue
        after = metadata(path)
        if after != entry:
            if not (entry['path'].endswith('.py') and '/scripts/' in entry['path']):
                raise ValueError(f'Unexpected evidence change: {entry["path"]}')
            changed.append({'path': entry['path'], 'before': entry, 'after': after})
        kept.append(after)
    # Delete only after both replacement artifacts and every retained baseline
    # entry pass validation. All tracked versions remain recoverable in Git.
    for path in current.values():
        path.unlink()
    manifest['files'] = kept
    manifest['active_revision'] = REVISION
    write(AUDIT / 'revision.json', {
        'revision': REVISION, 'created_at': datetime.now(timezone.utc).isoformat(),
        'prior_active_revision': prior.get('active_revision'),
        'prior_release_manifest_sha256': sha(AUDIT / 'prior-release-manifest.json'),
        'scope': 'User-requested relocation of all current summary XLSX/PDFs into summary/; obsolete copies removed; generators redirected; source study evidence retained.',
        'changes': changed, 'removed_files': inventory['files'],
        'validation': {mode: counts for mode, counts in validation['modes'].items()},
        'recovery': 'Previous tracked files are recoverable from the parent Git commit; no binary archive is retained in the repository.'})
    support = manifest.setdefault('supporting_files', [])
    for path in sorted(AUDIT.iterdir()):
        if path.is_file():
            support.append(metadata(path))
    write(ROOT / 'evidence-manifest.json', manifest)
    if old_artifacts():
        raise ValueError('Summary copies remain outside summary/')
    print(f'Removed {len(current)} inventoried artifacts. Recovery: parent Git commit.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action', choices=['prepare', 'validate', 'finish'])
    args = parser.parse_args()
    {'prepare': prepare, 'validate': validate, 'finish': finish}[args.action]()
