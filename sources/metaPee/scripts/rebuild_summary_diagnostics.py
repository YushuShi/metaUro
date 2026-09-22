"""Restyle saved diagnostic workbooks; never re-pool or rewrite evidence.

Run with --check to validate saved counts, or with no arguments to render.
Only current, canonical and active UI views are selected; archived interim
views are excluded. Scientific inputs are hashed before and after rendering.
"""
from pathlib import Path
import argparse, hashlib, inspect, json, os, shutil, subprocess, sys
ROOT = Path(__file__).resolve().parents[1]
METAAURO_ROOT = next((parent for parent in ROOT.parents
    if (parent/'summary_outputs.py').is_file() and (parent/'source_worker.py').is_file()), None)
if METAAURO_ROOT is not None and __name__ == '__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--check',action='store_true')
    args=parser.parse_args();sys.path.insert(0,str(METAAURO_ROOT))
    from summary_outputs import ensure_current
    for mode in ('yes','no'):
        ensure_current(METAAURO_ROOT,ROOT.name,mode,force=not args.check)
    print(f'{ROOT.name}: canonical summaries are current.',flush=True)
    raise SystemExit(0)
sys.path.insert(0, str(ROOT))
OUT = ROOT / 'outputs/diagnostic_square_log10_2026-09-17'
OUT.mkdir(parents=True, exist_ok=True)
os.environ.setdefault('MPLCONFIGDIR', str(OUT / 'matplotlib-cache'))
from openpyxl import load_workbook
from scripts import build_summary_plots as plots
from scripts import diagnostic_style as style


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def manifest_paths():
    names = subprocess.check_output(['rg', '--files', '--hidden', '-g', 'summary_manifest.json',
        '-g', '!outputs/**', '-g', '!.git/**', '-g', '!**/node_modules/**'], cwd=ROOT, text=True).splitlines()
    active_heart = ROOT / 'data/local_screening_active.json'
    heart_id = json.loads(active_heart.read_text())['view_id'] if active_heart.exists() else None
    active_gut = ROOT / 'combined/nonmalignant_extension_2026-09-10/active_interim_view.json'
    gut_view = json.loads(active_gut.read_text())['view_path'] if active_gut.exists() else None
    active_breath = ROOT / 'combined/nonmalignant_active_view.json'
    breath_view = json.loads(active_breath.read_text())['namespace'] if active_breath.exists() else None
    selected=[]
    for name in sorted(names):
        if 'local_screening_views/' in name and ('local_screening_views/' + str(heart_id) + '/') not in name: continue
        if '/interim_views/' in name and str(gut_view) + '/' not in name: continue
        if '/nonmalignant_extension_' in name and '_interim_' in name and str(breath_view) + '/' not in name: continue
        if name.startswith(('Plot/', 'combined/', 'data/subexposure_views/', 'data/local_screening_views/')):
            selected.append(ROOT/name)
    return selected


def mode_for(path, manifest):
    if manifest.get('include_subexposures') in ('yes','no'): return manifest['include_subexposures']
    for part in path.relative_to(ROOT).parts:
        if part in ('yes','no'): return part
        if part in ('subexposures_yes','subexposures_no'): return part.rsplit('_',1)[-1]
    return None


def read_rows(path):
    book=load_workbook(path, read_only=True, data_only=True)
    sheet=book.active; iterator=sheet.iter_rows(values_only=True); headers=next(iterator)
    rows=[]
    for values in iterator:
        row=dict(zip(headers, values))
        if not isinstance(row.get('number studies',row.get('number estimates')), (float,int)): continue
        row['category']=plots.exposure_category(row['Exposure'])
        rows.append(row)
    book.close()
    return rows


def plan():
    jobs=[]; inputs={}; mismatches=[]
    for path in manifest_paths():
        inputs[str(path.relative_to(ROOT))]=sha(path)
        manifest=json.loads(path.read_text()); mode=mode_for(path,manifest)
        for key, scope in manifest['scopes'].items():
            wanted=[('effect-heterogeneity','effect'),('egger-heterogeneity','egger')]
            if not any(scope['plots'].get(n,{}).get('available') for n,k in wanted): continue
            data_path=path.parent/'summary_rows.json'
            if data_path.exists():
                rows=json.loads(data_path.read_text())['combined']
            else:
                data_path=path.parent/scope['workbooks']['combined']['filename']
                rows=read_rows(data_path)
            inputs[str(data_path.relative_to(ROOT))]=sha(data_path)
            for name,kind in wanted:
                entry=scope['plots'].get(name,{})
                if not entry.get('available'): continue
                eligible=plots.heterogeneity_rows(rows,kind)
                if len(eligible)!=entry['exposure_count']:
                    mismatches.append({'manifest':str(path.relative_to(ROOT)),'scope':key,'plot':name,
                        'saved_count':entry['exposure_count'],'current_count':len(eligible)})
                    continue
                for locale,filename in entry.get('locales',{'en':entry['filename']}).items():
                    kwargs={}
                    if manifest.get('smoking_status_label'):
                        from scripts.smoking_status_summary import _translated
                        kwargs['smoking_status_label']=_translated('Smoking status',locale)+': '+_translated(manifest['smoking_status_label'],locale)
                        kwargs['subexposure_mode']=mode
                    if 'coverage_note' in inspect.signature(plots.heterogeneity_plot).parameters:
                        kwargs['coverage_note']=plots.screening_coverage_note(manifest.get('screening_coverage'),locale)
                    jobs.append(dict(path=path.parent/filename,rows=rows,disease=scope['label'],kind=kind,
                        locale=locale,mode=mode,kwargs=kwargs,count=len(eligible),
                        coordinates=[[r['Exposure'], r['eggers p-value'] if kind=='egger' else r['Pooled RR'],r['I^2 (%)']] for r in eligible]))
    return jobs,inputs,mismatches


def rebuild(check=False):
    jobs,inputs,mismatches=plan()
    if mismatches:
        (OUT/'count_mismatches.json').write_text(json.dumps(mismatches,indent=2)+'\n')
        print(json.dumps({'repo':ROOT.name,'mismatches':mismatches},indent=2),flush=True)
        return 2
    print(f'{ROOT}: {len(jobs)} diagnostic figures; saved counts match.',flush=True)
    if check: return 0
    audit_path=OUT/'render_audit.json'
    prior=json.loads(audit_path.read_text()) if audit_path.exists() else {}
    code_hash=hashlib.sha256(b''.join((ROOT/p).read_bytes() for p in
        ['scripts/build_summary_plots.py','scripts/plot_localization.py','scripts/diagnostic_style.py', 'scripts/rebuild_summary_diagnostics.py'])).hexdigest()
    audit={'style':style.STYLE_VERSION,'code_sha256':code_hash,'inputs_before':inputs,'outputs':{}}
    old=prior.get('outputs',{}) if prior.get('code_sha256')==code_hash and prior.get('inputs_before')==inputs else {}
    reused={}
    for index,job in enumerate(jobs,1):
        target=job['path']; relative=str(target.relative_to(ROOT)); png=target.with_suffix('.png')
        if relative in old and target.exists() and png.exists() and sha(target)==old[relative]['pdf_sha256'] and sha(png)==old[relative]['png_sha256']:
            audit['outputs'][relative]=old[relative]; continue
        identity=json.dumps([job['rows'],job['disease'],job['kind'],job['locale'],job['mode'],job['kwargs']],sort_keys=True,default=str)
        if identity in reused:
            source=reused[identity]
            target.parent.mkdir(parents=True,exist_ok=True)
            shutil.copy2(source,target);shutil.copy2(source.with_suffix('.png'),png)
        else:
            plots.SUBEXPOSURE_MODE=job['mode']
            count=plots.heterogeneity_plot(target,job['rows'],job['disease'],job['kind'],job['locale'],**job['kwargs'])
            if count!=job['count']: raise RuntimeError('Diagnostic eligibility changed while rendering')
            reused[identity]=target
        from PIL import Image
        with Image.open(png) as im:
            if im.width != im.height: raise RuntimeError('Non-square PNG: '+str(png))
            size=list(im.size)
        audit['outputs'][relative]={'pdf_sha256':sha(target),'png_sha256':sha(png),'png_pixels':size,
            'exposure_count':job['count'],'source_coordinates':job['coordinates'],'egger_transform':'log10' if job['kind']=='egger' else None}
        audit_path.write_text(json.dumps(audit,indent=2)+'\n')
        if index%20==0: print(f'{ROOT.name}: {index}/{len(jobs)}',flush=True)
    after={name:sha(ROOT/name) for name in inputs}
    if after!=inputs: raise RuntimeError('Scientific inputs changed during rendering')
    audit['inputs_after']=after
    audit_path.write_text(json.dumps(audit,indent=2)+'\n')
    print(f'{ROOT.name}: finished {len(jobs)} PDF/PNG pairs; scientific inputs unchanged.',flush=True)
    return 0


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--check',action='store_true')
    sys.exit(rebuild(parser.parse_args().check))
