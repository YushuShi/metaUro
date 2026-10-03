#!/usr/bin/env python3
"""Build cross-exposure summary workbooks and plots from saved UI caches."""

from __future__ import annotations

import argparse
import io
import json
import math
import re
import sys
import warnings
from collections import defaultdict
from contextlib import redirect_stdout
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import matplotlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from matplotlib.lines import Line2D

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib import patheffects  # noqa: E402
from matplotlib.gridspec import GridSpec  # noqa: E402
from matplotlib.ticker import FuncFormatter, LogLocator, NullFormatter  # noqa: E402
from matplotlib.transforms import Bbox  # noqa: E402


REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))
import meta_analysis  # noqa: E402
from scripts.summary_pi_style import draw_prediction_interval
from scripts.plot_localization import (  # noqa: E402
    category_label,
    disease_label,
    exposure_label,
    font_family,
    supported_languages,
    text as plot_text,
    validate_localizations,
)


COLUMNS = [
    "Exposure",
    "number estimates",
    "Pooled RR",
    "lower CI RR",
    "upper CI RR",
    "lower PI RR",
    "upper PI RR",
    "I^2 (%)",
    "eggers p-value",
    "total N",
    "total Cases",
]

from scripts.exposure_taxonomy import CATEGORY_MEMBERS, CATEGORY_COLORS, exposure_category as _shared_category


def safe_component(value: Any) -> str:
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]", "_", str(value).lower())).strip("_")


def _plain_log_tick(value: float, _position: int) -> str:
    """Format logarithmic ticks without math-text minus glyphs."""
    if not math.isfinite(value) or value <= 0:
        return ""
    return f"{value:g}"


def _use_plain_log_ticks(axis: Any) -> None:
    axis.xaxis.set_major_locator(LogLocator(base=10, subs=(1.0, 2.0, 5.0)))
    axis.xaxis.set_major_formatter(FuncFormatter(_plain_log_tick))
    axis.xaxis.set_minor_formatter(NullFormatter())


def exposure_category(slug: str) -> str:
    return _shared_category(slug, REPO_ROOT)


def _numeric(frame: pd.DataFrame, column: str) -> pd.Series:
    values = frame.get(column, pd.Series(np.nan, index=frame.index))
    return pd.to_numeric(
        values.astype(str).str.replace(",", "", regex=False).str.strip(),
        errors="coerce",
    )


def analysis_cases(frame: pd.DataFrame, disease: str) -> pd.Series:
    reported = _numeric(frame, "Cases")
    estimated = _numeric(frame, "Estimated Cases")
    sample_size = _numeric(frame, "Sample Size")
    baseline = meta_analysis.baseline_risk_for_disease(disease)
    derived = (
        np.floor(sample_size * float(baseline) + 0.5)
        if baseline is not None
        else pd.Series(np.nan, index=frame.index)
    )
    return reported.fillna(estimated).fillna(derived)


def eligible_effect_type(value: Any) -> bool:
    helper = getattr(meta_analysis, "is_eligible_effect_type", None)
    if callable(helper):
        return bool(helper(value))
    normalizer = getattr(meta_analysis, "normalize_effect_type", None)
    normalized = normalizer(value) if callable(normalizer) else str(value or "").upper()
    return str(normalized or "").upper() in {"OR", "RR", "IRR", "HR"}


def summarize_cache(path: Path, disease: str, exposure: str, dietary: bool) -> dict[str, Any] | None:
    payload = json.loads(path.read_text(encoding="utf-8"))
    frame = pd.DataFrame(payload.get("studies", []))
    if 'subexposure_pooling_selected' in frame.columns:
        frame = frame.loc[~frame['subexposure_pooling_selected'].eq(False)].copy()
    if frame.empty:
        return None
    if dietary:
        measurement = frame.get(
            "exposure_measurement_type", pd.Series("", index=frame.index)
        ).fillna("").astype(str).str.strip().str.lower()
        frame = frame.loc[measurement.eq("dietary_intake")].copy()
    if frame.empty:
        return None

    for column in (
        "Effect Size", "Lower CI", "Upper CI", "Cases", "Sample Size", "Estimated Cases"
    ):
        frame[column] = _numeric(frame, column)
    cases = analysis_cases(frame, disease)
    quality = frame.get("Quality Score", pd.Series("", index=frame.index)).fillna("").astype(str).str.strip().str.lower()
    effects = frame.get("Effect Type", pd.Series("", index=frame.index)).map(eligible_effect_type)
    valid = frame.loc[
        (frame["Effect Size"] > 0)
        & (frame["Lower CI"] > 0)
        & (frame["Upper CI"] > frame["Lower CI"])
        & frame["Sample Size"].gt(0)
        & frame["Sample Size"].mod(1).eq(0)
        & effects
        & quality.isin({"good", "moderate"})
    ].copy()
    from disease_hierarchy import select_frames
    valid, _ = select_frames(valid)
    if valid.empty:
        return None

    with warnings.catch_warnings(), redirect_stdout(io.StringIO()):
        warnings.simplefilter("ignore")
        parent_analysis = None
        if payload.get('disease_hierarchy_prebuilt'):
            from scripts.provisional_disease_scopes import analyze_parent
            parent_analysis = analyze_parent(valid, valid, disease, exposure, "Incidence")
            result = parent_analysis['result']
        else:
            result = meta_analysis.perform_meta_analysis(
                valid, disease, exposure, outcome="Incidence", exclude_meta=True,
                df_all=valid, generate_plots=False,
            )
    headline = result.get("headline") if isinstance(result, dict) else None
    if not isinstance(headline, dict) or headline.get("pooled_es") is None:
        return None
    if parent_analysis is not None:
        valid = pd.DataFrame(parent_analysis['pooled_rows'])
        if valid.empty:
            return None
        sample_sizes = _numeric(valid, 'Sample Size').fillna(0)
        reported_cases = _numeric(valid, 'Cases')
        baseline = parent_analysis['baseline'].get('p0')
        derived_cases = np.floor(sample_sizes * baseline + 0.5) if baseline is not None else 0
        case_counts = reported_cases.fillna(derived_cases).fillna(0)
        headline = headline.get('canonical_rr') or headline
    else:
        sample_sizes = valid["Sample Size"].fillna(0)
        case_counts = cases.loc[valid.index].fillna(0)
    return {
        "Exposure": exposure,
        "number estimates": int(len(valid)),
        "Pooled RR": float(headline["pooled_es"]),
        "lower CI RR": float(headline["ci_low"]),
        "upper CI RR": float(headline["ci_upp"]),
        "lower PI RR": headline.get("pi_low"),
        "upper PI RR": headline.get("pi_upp"),
        "I^2 (%)": round(float(headline.get("i2") or 0), 1),
        "eggers p-value": headline.get("eggers_p"),
        "total N": int(sample_sizes.sum()),
        "total Cases": int(case_counts.sum()),
        "category": exposure_category(exposure),
    }


def inventory(cache_root: Path) -> dict[str, list[tuple[Path, str]]]:
    grouped: dict[str, list[tuple[Path, str]]] = defaultdict(list)
    for path in sorted(cache_root.glob("*/*_incidence_true_all.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        # Authoritative revision runs deliberately persist zero-inclusion
        # caches so a searched scope cannot fall back to stale evidence.  They
        # are audit state, not visible summary evidence, and therefore must not
        # enter localization validation or plot/workbook inventories.
        studies = payload.get("studies")
        if not isinstance(studies, list) or not studies:
            continue
        disease = str(payload.get("cancer") or "").strip()
        if disease:
            grouped[disease].append((path, path.parent.name))
    return dict(grouped)


def workbook(path: Path, rows: list[dict[str, Any]]) -> None:
    export = pd.DataFrame(rows, columns=COLUMNS)
    export.to_excel(path, index=False)
    book = load_workbook(path)
    if SUBEXPOSURE_MODE is not None:
        info = book.create_sheet('Analysis mode')
        info.append(['Include subexposures', SUBEXPOSURE_MODE.title()])
        info.append(['Scope', 'Transitive child evidence with overlap protection' if SUBEXPOSURE_MODE == 'yes' else 'Direct exposure evidence only'])

    sheet = book.active
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30
    widths = [28, 15, 13, 14, 14, 14, 14, 12, 16, 15, 15]
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1A1A2E")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for index, cell in enumerate(row, start=1):
            if index in {2, 10, 11}:
                cell.number_format = "#,##0"
            elif index == 8:
                cell.number_format = "0.0"
            elif index == 9:
                cell.number_format = "0.0000"
            elif index > 2:
                cell.number_format = "0.00"
    book.save(path)


def _ordered_plot_rows(rows: list[dict[str, Any]], direction: str) -> list[dict[str, Any]]:
    eligible = [
        row for row in rows
        if row["number estimates"] > 1
        and ((row["Pooled RR"] < 1) if direction == "protective" else (row["Pooled RR"] >= 1))
    ]
    order = list(CATEGORY_MEMBERS) + ["Other"]
    return sorted(
        eligible,
        key=lambda row: (
            order.index(row["category"]) if row["category"] in order else len(order),
            row["Pooled RR"],
        ),
    )


def _forest_caption(
    direction: str,
    disease: str,
    locale: str,
    interval_truncated: bool = False,
    has_significant: bool = False,
) -> str:
    disease_text = disease_label(disease, locale)
    association_key = "inverse_caption" if direction == "protective" else "positive_caption"
    parts = [
        plot_text(association_key, locale, disease=disease_text),
        plot_text("pooled_marker_caption", locale),
        *([plot_text("significant_caption", locale)] if has_significant else []),
        plot_text("nonsignificant_caption", locale),
        plot_text("prediction_interval_caption", locale),
        plot_text("estimate_overlap_caption", locale),
    ]
    if interval_truncated:
        parts.append(plot_text("interval_truncation_caption", locale))
    return " | ".join(parts)


def _draw_bounded_interval(
    axis: Any,
    low: Any,
    high: Any,
    y: float,
    x_min: float,
    x_max: float,
    *,
    color: str,
    linewidth: float,
    linestyle: str = "-",
    alpha: float = 1.0,
    zorder: float = 2.0,
) -> bool:
    """Draw an interval within a log axis and mark either truncated end."""
    try:
        low_value = float(low)
        high_value = float(high)
    except (TypeError, ValueError):
        return False
    if math.isnan(low_value) or math.isnan(high_value):
        return False
    extends_left = not math.isfinite(low_value) or low_value <= 0 or low_value < x_min
    extends_right = not math.isfinite(high_value) or high_value > x_max
    visible_low = x_min if extends_left else low_value
    visible_high = x_max if extends_right else high_value
    if visible_low <= visible_high:
        axis.plot(
            [visible_low, visible_high],
            [y, y],
            color=color,
            linewidth=linewidth,
            linestyle=linestyle,
            alpha=alpha,
            solid_capstyle="round",
            zorder=zorder,
        )
    marker_size = 22 if linewidth >= 1 else 12
    if extends_left:
        axis.scatter(
            x_min,
            y,
            s=marker_size,
            marker="<",
            color=color,
            alpha=alpha,
            linewidths=0,
            clip_on=False,
            zorder=zorder + 0.1,
        )
    if extends_right:
        axis.scatter(
            x_max,
            y,
            s=marker_size,
            marker=">",
            color=color,
            alpha=alpha,
            linewidths=0,
            clip_on=False,
            zorder=zorder + 0.1,
        )
    return extends_left or extends_right


def forest_plot(
    path: Path,
    rows: list[dict[str, Any]],
    disease: str,
    direction: str,
    dietary: bool,
    locale: str = "en",
) -> int:
    rows = _ordered_plot_rows(rows, direction)
    if not rows:
        path.unlink(missing_ok=True)
        return 0
    display_rows: list[tuple[str, dict[str, Any] | None]] = []
    for category in dict.fromkeys(row["category"] for row in rows):
        display_rows.append((category, None))
        display_rows.extend((category, row) for row in rows if row["category"] == category)

    height = max(5.8, 0.38 * len(display_rows) + 3.0)
    family = font_family(locale)
    with plt.rc_context({
        "font.family": family,
        "axes.unicode_minus": False,
        "mathtext.fontset": "dejavusans",
    }):
        fig = plt.figure(figsize=(13.5, height))
        grid = GridSpec(
            1,
            2,
            figure=fig,
            width_ratios=[2.25, 1.0],
            wspace=0.02,
            left=0.025,
            right=0.985,
            top=0.90,
            bottom=min(0.24, max(0.12, 1.25 / height)),
        )
        table_ax = fig.add_subplot(grid[0, 0])
        forest_ax = fig.add_subplot(grid[0, 1])
        table_ax.set_xlim(0, 1)
        table_ax.set_ylim(len(display_rows) + 0.7, -1.2)
        forest_ax.set_ylim(len(display_rows) + 0.7, -1.2)
        table_ax.axis("off")
        forest_ax.set_xscale("log")
        _use_plain_log_ticks(forest_ax)
        low = min(float(row["lower CI RR"]) for row in rows)
        high = max(float(row["upper CI RR"]) for row in rows)
        x_min = max(0.05, 10 ** math.floor(math.log10(low)))
        x_max = min(20.0, 10 ** math.ceil(math.log10(high)))
        if x_min >= 1:
            x_min = 0.5
        if x_max <= 1:
            x_max = 2.0
        forest_ax.set_xlim(x_min, x_max)
        forest_ax.axvline(1, color="#1A1A2E", linestyle="--", linewidth=1)
        forest_ax.grid(axis="x", color="#D7DEE5", linestyle=":", linewidth=0.6)
        forest_ax.set_yticks([])
        forest_ax.set_xlabel(plot_text("pooled_rr", locale), fontweight="bold")
        for spine in ("top", "left", "right"):
            forest_ax.spines[spine].set_visible(False)

        headers = [
            (0.01, plot_text("exposure", locale)),
            (0.49, plot_text("number_of_estimates", locale)),
            (0.62, plot_text("sample_size", locale)),
            (0.75, plot_text("cases", locale)),
            (0.98, plot_text("pooled_rr_ci", locale)),
        ]
        for x, label in headers:
            table_ax.text(
                x,
                -0.55,
                label,
                ha="right" if x > 0.8 else ("left" if x < 0.1 else "center"),
                va="center",
                color="white",
                fontweight="bold",
                fontsize=8.4,
            )
        table_ax.axhspan(-1.05, -0.05, color="#1A1A2E")
        forest_ax.axhspan(-1.05, -0.05, color="#1A1A2E")
        forest_ax.text(
            1,
            -0.55,
            plot_text("effect_size", locale),
            ha="center",
            va="center",
            color="white",
            fontweight="bold",
            fontsize=8.4,
        )

        interval_truncated = False
        for y, (category, row) in enumerate(display_rows, start=1):
            color = CATEGORY_COLORS.get(category, CATEGORY_COLORS["Other"])
            if row is None:
                table_ax.axhspan(y - 0.48, y + 0.48, color=color, alpha=0.16)
                forest_ax.axhspan(y - 0.48, y + 0.48, color=color, alpha=0.16)
                table_ax.text(
                    0.01,
                    y,
                    category_label(category, locale).upper(),
                    color=color,
                    fontweight="bold",
                    fontstyle="italic",
                    fontsize=8.2,
                    va="center",
                )
                continue
            if y % 2 == 0:
                table_ax.axhspan(y - 0.48, y + 0.48, color="#F5F7FA")
                forest_ax.axhspan(y - 0.48, y + 0.48, color="#F5F7FA")
            table_ax.text(
                0.02,
                y,
                exposure_label(REPO_ROOT, row["Exposure"], locale),
                color=color,
                fontweight="bold",
                fontsize=8.3,
                va="center",
            )
            table_ax.text(0.49, y, f'{row["number estimates"]:,}', ha="center", va="center", fontsize=8.0, color=color)
            table_ax.text(0.62, y, f'{row["total N"]:,}', ha="center", va="center", fontsize=8.0, color=color)
            table_ax.text(0.75, y, f'{row["total Cases"]:,}', ha="center", va="center", fontsize=8.0, color=color)
            table_ax.text(0.98, y, f'{row["Pooled RR"]:.2f} ({row["lower CI RR"]:.2f}-{row["upper CI RR"]:.2f})', ha="right", va="center", fontsize=8.0, color=color, fontstyle="italic")
            interval_truncated = draw_prediction_interval(
                forest_ax, row.get("lower PI RR"), row.get("upper PI RR"),
                y, x_min, x_max) or interval_truncated
            interval_truncated = _draw_bounded_interval(
                forest_ax,
                row["lower CI RR"],
                row["upper CI RR"],
                y,
                x_min,
                x_max,
                color=color,
                linewidth=1.5,
                zorder=2.2,
            ) or interval_truncated
            significant = row["upper CI RR"] < 1 or row["lower CI RR"] > 1
            marker_size = 24 + row["number estimates"] * 4
            if significant:
                forest_ax.scatter(row["Pooled RR"], y, s=marker_size * 1.7, marker="D", color="white", alpha=0.65, linewidths=0, zorder=3)
            forest_ax.scatter(
                row["Pooled RR"],
                y,
                s=marker_size,
                marker="D",
                color=color,
                alpha=1.0 if significant else 0.38,
                edgecolors="#CC0000" if significant else "none",
                linewidths=1.1 if significant else 0,
                zorder=4,
            )

        disease_text = disease_label(disease, locale)
        title_key = "protective_title" if direction == "protective" else "harmful_title"
        title = plot_text(title_key, locale, disease=disease_text)
        if dietary:
            title += "\n" + plot_text("dietary_subtitle", locale)
        fig.suptitle(title, fontsize=13, fontweight="bold", color="#1A1A2E", y=0.96)

        categories_present = list(dict.fromkeys(row["category"] for row in rows))
        legend_handles = [
            Line2D(
                [0],
                [0],
                marker="s",
                linestyle="none",
                markerfacecolor=CATEGORY_COLORS.get(category, CATEGORY_COLORS["Other"]),
                markeredgecolor="none",
                markersize=7,
                label=category_label(category, locale),
            )
            for category in categories_present
        ]
        fig.legend(
            handles=legend_handles,
            title=plot_text("exposure_group", locale),
            loc="lower center",
            bbox_to_anchor=(0.5, 0.052),
            ncol=min(5, max(1, len(legend_handles))),
            frameon=False,
            fontsize=7.3,
            title_fontsize=8.0,
            columnspacing=1.1,
            handletextpad=0.35,
        )
        fig.text(
            0.02,
            0.012,
            _forest_caption(direction, disease, locale, interval_truncated,
                            has_significant=any((row["upper CI RR"] < 1 or row["lower CI RR"] > 1)
                                                and x_min <= row["Pooled RR"] <= x_max for row in rows)),
            ha="left",
            va="bottom",
            fontsize=6.8,
            color="#607D8B",
            fontstyle="italic",
            wrap=True,
        )
        path.parent.mkdir(parents=True, exist_ok=True)
        if SUBEXPOSURE_MODE is not None:
            fig.text(0.99, 0.005, 'Include subexposures: ' + SUBEXPOSURE_MODE.title(), ha='right', va='bottom', fontsize=7, color='#607D8B')
        fig.savefig(path, format="pdf", bbox_inches="tight", metadata={"CreationDate": None})
        plt.close(fig)
    return len(rows)


def _intersection_area(first: Bbox, second: Bbox) -> float:
    width = max(0.0, min(first.x1, second.x1) - max(first.x0, second.x0))
    height = max(0.0, min(first.y1, second.y1) - max(first.y0, second.y0))
    return width * height


def _close_label_candidates(
    point: tuple[float, float],
    width: float,
    height: float,
    marker_size: float,
    dpi: float,
) -> list[tuple[float, float, str, str, Bbox, float]]:
    """Return compact, line-free positions surrounding a scatter marker."""

    point_x, point_y = point
    point_to_pixel = dpi / 72.0
    marker_radius = math.sqrt(max(marker_size, 1.0)) / 2.0
    base_gap = marker_radius + 2.4
    angles = (
        0, 180, 90, 270, 45, 135, 315, 225,
        22.5, 157.5, 337.5, 202.5, 67.5, 112.5, 292.5, 247.5,
    )
    candidates = []
    distance = base_gap
    for angle in angles:
        radians = math.radians(angle)
        dx = distance * math.cos(radians)
        dy = distance * math.sin(radians)
        horizontal_alignment = "left" if dx > 1 else "right" if dx < -1 else "center"
        vertical_alignment = "bottom" if dy > 1 else "top" if dy < -1 else "center"
        anchor_x = point_x + dx * point_to_pixel
        anchor_y = point_y + dy * point_to_pixel
        if horizontal_alignment == "left":
            x0 = anchor_x
        elif horizontal_alignment == "right":
            x0 = anchor_x - width
        else:
            x0 = anchor_x - width / 2.0
        if vertical_alignment == "bottom":
            y0 = anchor_y
        elif vertical_alignment == "top":
            y0 = anchor_y - height
        else:
            y0 = anchor_y - height / 2.0
        candidates.append(
            (
                dx,
                dy,
                horizontal_alignment,
                vertical_alignment,
                Bbox.from_bounds(x0, y0, width, height),
                distance,
            )
        )
    return candidates


def _place_close_labels(fig: Any, ax: Any, labels: list[dict[str, Any]]) -> None:
    """Choose less crowded sides while keeping every label at its marker edge."""

    if not labels:
        return
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    axes_box = ax.get_window_extent(renderer=renderer)
    occupied = []
    for label in labels:
        annotation = label["annotation"]
        bounds = annotation.get_window_extent(renderer=renderer)
        point = tuple(ax.transData.transform((label["x"], label["y"])))
        candidates = _close_label_candidates(
            point,
            max(bounds.width, 1.0),
            max(bounds.height, 1.0),
            label["marker_size"],
            fig.dpi,
        )
        # Favor a clear side without increasing the marker-to-label gap.
        placement = min(
            candidates,
            key=lambda candidate: (
                max(0.0, candidate[4].width * candidate[4].height
                    - _intersection_area(candidate[4], axes_box)),
                sum(_intersection_area(candidate[4], box) for box in occupied),
            ),
        )
        occupied.append(placement[4].padded(1.5))
        dx, dy, horizontal_alignment, vertical_alignment, _, _ = placement
        annotation.set_position((dx, dy))
        annotation.set_ha(horizontal_alignment)
        annotation.set_va(vertical_alignment)

    # Keep the nearest placements when they fit; resolve crowded groups together.
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    boxes = [label["annotation"].get_window_extent(renderer) for label in labels]
    points = ax.transData.transform([(label["x"], label["y"]) for label in labels])
    crowded = any(box.overlaps(other) for i, box in enumerate(boxes) for other in boxes[i + 1:])
    crowded = crowded or any(box.contains(*point) for box in boxes for point in points)
    crowded = crowded or any(box.x0 < axes_box.x0 or box.x1 > axes_box.x1
                             or box.y0 < axes_box.y0 or box.y1 > axes_box.y1 for box in boxes)
    if crowded:
        from scripts.diagnostic_labels import place_labels
        for label in labels:
            label.setdefault("text", label["annotation"].get_text())
            label.setdefault("color", label["annotation"].get_color())
        place_labels(fig, ax, labels)


def heterogeneity_plot(
    path: Path,
    rows: list[dict[str, Any]],
    disease: str,
    kind: str,
    locale: str = "en",
) -> int:
    eligible = [
        row
        for row in rows
        if row["number estimates"] >= (10 if kind == "egger" else 3)
        and row.get("I^2 (%)") is not None
        and math.isfinite(float(row["I^2 (%)"]))
    ]
    if kind == "egger":
        eligible = [row for row in eligible if row.get("eggers p-value") not in (None, 0) and 0 < float(row["eggers p-value"]) <= 1]
    if len(eligible) < 2:
        path.unlink(missing_ok=True)
        return 0

    family = font_family(locale)
    with plt.rc_context({
        "font.family": family,
        "axes.unicode_minus": False,
        "mathtext.fontset": "dejavusans",
    }):
        fig, ax = plt.subplots(figsize=(11.7, 8.3))
        x_values = [
            math.log(float(row["eggers p-value"])) if kind == "egger" else float(row["Pooled RR"])
            for row in eligible
        ]
        labels = []
        for index, row in enumerate(eligible):
            category = row["category"]
            color = CATEGORY_COLORS.get(category, CATEGORY_COLORS["Other"])
            x = x_values[index]
            y = float(row["I^2 (%)"])
            marker_size = 45 + row["number estimates"] * 6
            ax.scatter(
                x,
                y,
                s=marker_size,
                color=color,
                alpha=0.86,
                edgecolors="white",
                linewidths=0.9,
                zorder=3,
            )
            labels.append({
                "x": x,
                "y": y,
                "marker_size": marker_size,
                "color": color,
                "text": exposure_label(REPO_ROOT, row["Exposure"], locale),
            })

        # Keep valid I² = 0 markers and their labels clear of the bottom spine.
        # The visible tick scale remains the meaningful 0-100% range.
        ax.set_ylim(-6, 102)
        ax.set_yticks(range(0, 101, 20))
        ax.set_ylabel(plot_text("heterogeneity_y", locale), fontweight="bold", fontsize=13)
        ax.grid(color="#D7DEE5", linestyle=":", linewidth=0.7)
        disease_text = disease_label(disease, locale)
        if kind == "egger":
            ax.axvline(math.log(0.05), color="#1A1A2E", linestyle="--", linewidth=1)
            x_min = min(x_values)
            x_max = max(x_values)
            padding = max(0.35, (x_max - x_min) * 0.28)
            ax.set_xlim(x_min - padding, x_max + padding)
            ax.set_xlabel(plot_text("egger_x", locale), fontweight="bold", fontsize=13)
            title = plot_text("egger_heterogeneity_title", locale, disease=disease_text)
            caption = plot_text("egger_diagnostic_caption", locale)
        else:
            ax.axvline(1, color="#1A1A2E", linestyle="--", linewidth=1)
            ax.set_xscale("log")
            _use_plain_log_ticks(ax)
            ax.set_xlim(
                min(min(x_values) * 0.72, 0.82),
                max(max(x_values) * 1.38, 1.22),
            )
            ax.set_xlabel(plot_text("effect_heterogeneity_x", locale), fontweight="bold", fontsize=13)
            title = plot_text("effect_heterogeneity_title", locale, disease=disease_text)
            caption = plot_text("effect_diagnostic_caption", locale)
        ax.set_title(title, fontweight="bold", fontsize=19, color="#1A1A2E", pad=16)
        ax.tick_params(axis="both", labelsize=13)
        for tick in (*ax.get_xticklabels(), *ax.get_yticklabels()):
            tick.set_fontweight("bold")
        for spine in ("top", "right"):
            ax.spines[spine].set_visible(False)
        fig.text(0.5, 0.022, caption, ha="center", va="bottom", fontsize=10.0, color="#607D8B", fontstyle="italic", wrap=True)
        fig.subplots_adjust(left=0.10, right=0.98, top=0.89, bottom=0.15)
        font_size = 11.0
        for label in labels:
            label["annotation"] = ax.annotate(
                label["text"],
                xy=(label["x"], label["y"]),
                xycoords="data",
                xytext=(0, 0),
                textcoords="offset points",
                ha="center",
                va="center",
                fontsize=font_size,
                fontweight="bold",
                color=label["color"],
                path_effects=[patheffects.withStroke(linewidth=2.5, foreground="white")],
                annotation_clip=True,
                zorder=4,
            )
        _place_close_labels(fig, ax, labels)
        path.parent.mkdir(parents=True, exist_ok=True)
        if SUBEXPOSURE_MODE is not None:
            fig.text(0.99, 0.005, 'Include subexposures: ' + SUBEXPOSURE_MODE.title(), ha='right', va='bottom', fontsize=9, color='#607D8B')
        fig.savefig(path, format="pdf", bbox_inches="tight", metadata={"CreationDate": None})
        plt.close(fig)
    return len(eligible)


def _localized_plot_path(plot_root: Path, filename: str, locale: str) -> Path:
    return plot_root / filename if locale == "en" else plot_root / "locales" / locale / filename


SUBEXPOSURE_MODE = None


def build(
    cache_root: Path,
    plot_root: Path,
    languages: tuple[str, ...] | None = None,
) -> dict[str, Any]:
    global SUBEXPOSURE_MODE
    mode_path = cache_root / '_subexposure_manifest.json'
    SUBEXPOSURE_MODE = json.loads(mode_path.read_text()).get('include_subexposures') if mode_path.exists() else None
    plot_root.mkdir(parents=True, exist_ok=True)
    cache_inventory = inventory(cache_root)
    languages = tuple(languages or supported_languages(REPO_ROOT))
    if "en" not in languages:
        languages = ("en", *languages)
    unsupported = set(languages).difference(supported_languages(REPO_ROOT))
    if unsupported:
        raise ValueError(f"Unsupported plot languages: {sorted(unsupported)}")
    all_exposures = {exposure for caches in cache_inventory.values() for _, exposure in caches}
    all_categories = {exposure_category(exposure) for exposure in all_exposures}
    validate_localizations(
        REPO_ROOT,
        exposures=all_exposures,
        diseases=cache_inventory,
        categories=all_categories,
        locales=languages,
    )
    manifest: dict[str, Any] = {
        "schema_version": 2,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "eligibility_note": "Good or Moderate JBI and an eligible ratio effect. Counts are extraction estimates, including separate sex or measurement strata; forest summaries require at least two estimates per exposure. N and case totals sum estimates and may overlap across strata.",
        "languages": list(languages),
        "scopes": {},
    }
    for disease, caches in sorted(cache_inventory.items()):
        scope_slug = safe_component(disease)
        datasets: dict[str, list[dict[str, Any]]] = {}
        workbook_entries = {}
        for name, dietary in (("combined", False), ("dietary", True)):
            rows = []
            for path, exposure in caches:
                try:
                    row = summarize_cache(path, disease, exposure, dietary)
                except Exception as exc:
                    print(f"Skipped {path}: {exc}", file=sys.stderr)
                    continue
                if row:
                    rows.append(row)
            rows.sort(key=lambda row: row["Pooled RR"])
            datasets[name] = rows
            filename = f"exposures_meta_analysis_{scope_slug}_{name}.xlsx"
            workbook(plot_root / filename, rows)
            workbook_entries[name] = {"filename": filename, "available": bool(rows), "exposure_count": len(rows)}

        plot_specs = {
            "forest-protective": ("combined", "protective", False),
            "forest-harmful": ("combined", "harmful", False),
            "forest-protective-dietary": ("dietary", "protective", True),
            "forest-harmful-dietary": ("dietary", "harmful", True),
        }
        plots = {}
        for name, (dataset, direction, dietary) in plot_specs.items():
            suffix = f"{scope_slug}_dietary" if dietary else scope_slug
            filename = f"forest_{direction}_{suffix}.pdf"
            localized_files = {}
            count = 0
            for locale in languages:
                target = _localized_plot_path(plot_root, filename, locale)
                locale_count = forest_plot(
                    target,
                    datasets[dataset],
                    disease,
                    direction,
                    dietary,
                    locale,
                )
                if locale_count:
                    localized_files[locale] = str(target.relative_to(plot_root))
                if locale == "en":
                    count = locale_count
            plots[name] = {
                "filename": filename,
                "available": count > 0,
                "exposure_count": count,
                "locales": localized_files,
            }
        for name, kind in (("effect-heterogeneity", "effect"), ("egger-heterogeneity", "egger")):
            filename = f"plot_{'es' if kind == 'effect' else 'eggers'}_vs_heterogeneity_{scope_slug}.pdf"
            localized_files = {}
            count = 0
            for locale in languages:
                target = _localized_plot_path(plot_root, filename, locale)
                locale_count = heterogeneity_plot(
                    target,
                    datasets["combined"],
                    disease,
                    kind,
                    locale,
                )
                if locale_count:
                    localized_files[locale] = str(target.relative_to(plot_root))
                if locale == "en":
                    count = locale_count
            plots[name] = {
                "filename": filename,
                "available": count > 0,
                "exposure_count": count,
                "locales": localized_files,
            }

        manifest["scopes"][scope_slug] = {
            "label": disease,
            "cache_count": len(caches),
            "workbooks": workbook_entries,
            "plots": plots,
        }
    (plot_root / "summary_manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
    )
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cache-root", type=Path, default=REPO_ROOT / "Cached_results")
    parser.add_argument("--plot-root", type=Path, default=REPO_ROOT / "Plot")
    parser.add_argument(
        "--languages",
        nargs="+",
        choices=supported_languages(REPO_ROOT),
        help="Generate only these supported language variants (default: all).",
    )
    args = parser.parse_args()
    manifest = build(
        args.cache_root,
        args.plot_root,
        tuple(args.languages) if args.languages else None,
    )
    print(json.dumps({
        "scopes": len(manifest["scopes"]),
        "combined_exposures": sum(scope["workbooks"]["combined"]["exposure_count"] for scope in manifest["scopes"].values()),
        "available_plots": sum(plot["available"] for scope in manifest["scopes"].values() for plot in scope["plots"].values()),
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
