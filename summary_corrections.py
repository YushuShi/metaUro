"""Compatibility accessors for summaries now published only in summary/."""
from pathlib import Path
import json

from openpyxl import load_workbook

from summary_outputs import cache_dir, ensure_current, metauro_root, summary_dir, summary_mode

PLOT_LANGUAGES = frozenset({'en', 'zh-CN', 'zh-TW', 'nl', 'ko'})


def corrected_workbook(repo_root, mode, scope_key):
    folder = ensure_current(repo_root, 'metaPee', mode)
    manifest = json.loads((folder / 'summary_manifest.json').read_text())
    entry = manifest['scopes'][scope_key]['workbooks']['combined']
    if not entry['available']:
        raise FileNotFoundError(scope_key)
    return folder / entry['filename']


def corrected_rows(source_root, mode, scope_key, *, workbook=None):
    repo = metauro_root(source_root)
    if repo is None:
        raise ValueError('Expected a source inside metaUro')
    path = corrected_workbook(repo, mode, scope_key)
    book = workbook or load_workbook(path)
    sheet = book.active
    headers = [cell.value for cell in sheet[1]]
    rows = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
    manifest = json.loads((summary_dir(repo, 'metaPee', mode) / 'summary_manifest.json').read_text())
    cache = cache_dir(repo, 'metaPee', mode)
    corrected = 0
    for row in rows:
        path = cache / row['Exposure'] / f'{scope_key}_incidence_true_all.json'
        if path.is_file() and json.loads(path.read_text()).get('disease_hierarchy_prebuilt') is True:
            corrected += 1
        row['category'] = 'Other'
    return book, rows, manifest['scopes'][scope_key]['label'], corrected, path


def corrected_effect_heterogeneity_plot(repo_root, mode, scope_key, language='en'):
    folder = ensure_current(repo_root, 'metaPee', mode)
    manifest = json.loads((folder / 'summary_manifest.json').read_text())
    entry = manifest['scopes'][scope_key]['plots']['effect-heterogeneity']
    if not entry['available']:
        raise FileNotFoundError(scope_key)
    return folder / entry.get('locales', {}).get(language, entry['filename'])
