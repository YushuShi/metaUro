#!/usr/bin/env python3
"""Build optional cohort-selected summaries without changing source snapshots."""
from __future__ import annotations

import argparse
import importlib.util
import json
import os
from pathlib import Path
import sys
sys.dont_write_bytecode=True

ROOT = Path(__file__).resolve().parents[1]


def build(source, mode, languages=None,keep_latest_cohort=True,include_preprints=False):
    os.environ['METAURO_KEEP_LATEST_COHORT']='yes' if keep_latest_cohort else 'no'
    os.environ['METAURO_INCLUDE_PREPRINTS']='yes' if include_preprints else 'no'
    os.environ['MPLBACKEND'] = 'Agg'
    sys.path.insert(0, str(ROOT))
    from automation.sync_sources import assembled_source
    from summary_outputs import cache_dir, summary_dir
    source_root = assembled_source(ROOT, source)
    os.environ['METAURO_SUMMARY_SOURCE_ROOT'] = str(source_root)
    sys.path[:0] = [str(ROOT / 'runtime_code' / source), str(source_root)]
    import meta_analysis
    import analysis_settings_backend as backend
    import pandas as pd

    spec = importlib.util.spec_from_file_location('metauro_cohort_summary_builder',
        ROOT / 'sources' / source / 'scripts' / 'build_summary_plots.py')
    renderer = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = renderer
    spec.loader.exec_module(renderer)
    audits = []
    source_errors = []
    original = renderer.summarize_cache

    def summarize(path, disease, exposure, dietary):
        try:
            with backend.latest_cohort_pooling(keep_latest_cohort) as context:
                row = original(path, disease, exposure, dietary)
        except Exception as error:
            if include_preprints:source_errors.append((str(path),type(error).__name__))
            raise
        audits.append({'disease':disease,'exposure':exposure,'dataset':'dietary' if dietary else 'combined',**context.get('audit',{})})
        if row is None:
            return None
        selected = pd.DataFrame(context.get('rows', []))
        if selected.empty:
            return None
        cases = renderer.analysis_cases(selected, disease)
        row['number studies'] = len(selected)
        row['total N'] = int(pd.to_numeric(selected['Sample Size'], errors='coerce').fillna(0).sum())
        row['total Cases'] = int(cases.sum()) if cases.notna().all() else None
        if include_preprints:row['Preprint sources']=sum(item.get('preprint_derived') is True for item in context.get('rows',[]))
        return row

    renderer.summarize_cache = summarize
    from matplotlib.figure import Figure
    original_save = Figure.savefig
    active_preprint_exposures=set()
    if include_preprints:
        original_label=renderer.exposure_label
        def source_exposure_label(repo,exposure,locale):
            return original_label(repo,exposure,locale)+(' [preprint]' if exposure in active_preprint_exposures else '')
        renderer.exposure_label=source_exposure_label
        def source_plot(original_plot):
            def draw(path,rows,*args,**kwargs):
                nonlocal active_preprint_exposures
                previous=active_preprint_exposures
                active_preprint_exposures={row['Exposure'] for row in rows if row.get('Preprint sources',0)>0}
                try:return original_plot(path,rows,*args,**kwargs)
                finally:active_preprint_exposures=previous
            return draw
        renderer.forest_plot=source_plot(renderer.forest_plot)
        renderer.heterogeneity_plot=source_plot(renderer.heterogeneity_plot)

    def save_latest_figure(figure, *args, **kwargs):
        if include_preprints:
            if not any(text.get_text()=='Reviewed preprint recoveries included' for text in figure.texts):
                figure.text(.5,1.005,'Reviewed preprint recoveries included',ha='center',va='bottom',fontsize=8,color='#000080')
            for axis in figure.axes:
                for text in axis.texts:
                    if '[preprint]' in text.get_text():text.set_color('#000080')
        if keep_latest_cohort and not any(text.get_text() == 'Latest publication per cohort' for text in figure.texts):
            figure.text(.01, .005, 'Latest publication per cohort', ha='left', va='bottom', fontsize=7, color='#607D8B')
        return original_save(figure, *args, **kwargs)

    Figure.savefig = save_latest_figure
    original_finalize = renderer.finalize_build

    def finalize(source_root, folder, manifest):
        if include_preprints and source_errors:
            raise RuntimeError('Optional summaries were not published because native source calculations failed: '+str(source_errors[:3]))
        import hashlib
        from cohort_summary_support import bind_cohort_manifest
        if keep_latest_cohort:bind_cohort_manifest(manifest, ROOT)
        if include_preprints:
            from preprint_ui import preprint_manifest_stamp
            manifest['include_preprints']=True
            manifest['preprint_recovery_manifest_sha256']=preprint_manifest_stamp(ROOT)
        manifest['cohort_selection_audits'] = audits
        manifest['cohort_selection_audits_by_source'] = {source: audits}
        manifest['keep_latest_cohort'] = bool(keep_latest_cohort)
        manifest['cohort_policy'] = 'Latest definitely dated publication per identified cohort; numbered cohorts, ties and unknowns retained separately.'
        return original_finalize(source_root, folder, manifest)

    renderer.finalize_build = finalize
    original_workbook = renderer.workbook

    def workbook(path, rows, *args, **kwargs):
        original_workbook(path, rows, *args, **kwargs)
        if not keep_latest_cohort:return
        from openpyxl import load_workbook
        book = load_workbook(path)
        sheet = book.create_sheet('Cohort selection')
        sheet.append(['Keep latest publication per cohort', 'Yes'])
        sheet.append(['Unknown cohort/date or tied publication dates', 'Retained'])
        sheet.append(['Recorded source evidence', 'All records preserved; include-all is the default.'])
        sheet.append(['Disease','Exposure','Earlier PMID','Retained PMID','Cohorts'])
        disease = args[0] if args else kwargs.get('disease')
        dataset = 'dietary' if '_dietary' in Path(path).stem else 'combined'
        for audit in audits:
            if disease and audit['disease'] != disease or audit['dataset'] != dataset: continue
            for item in audit.get('excluded', []):
                sheet.append([audit['disease'], audit['exposure'], item['pmid'], item['retained_pmid'], '; '.join(item['cohort_ids'])])
        sheet.column_dimensions['A'].width = 48
        sheet.column_dimensions['B'].width = 80
        book.save(path)

    renderer.workbook = workbook
    if not include_preprints:
        return renderer.build(cache_dir(ROOT, source, mode), summary_dir(ROOT, source, mode), languages)
    import tempfile,shutil,hashlib
    from preprint_ui import compose_preprint_view,applicable_preprint_contexts,write_preprint_summary_seal
    from automation.sync_sources import _effective_file,load_overlay
    contexts=applicable_preprint_contexts(ROOT,{'source_repo':source,'include_subexposures':mode})
    with tempfile.TemporaryDirectory(prefix='metauro-preprint-caches-') as temporary:
        stage=Path(temporary)/'Cached_results'
        shutil.copytree(cache_dir(ROOT,source,mode),stage)
        for path in stage.glob('*/*_incidence_true_all.json'):
            original_cache_path=cache_dir(ROOT,source,mode)/path.relative_to(stage)
            relative=str(original_cache_path.resolve().relative_to(source_root.resolve()))
            durable=_effective_file(ROOT,source,relative,load_overlay(ROOT,source))
            if durable is None:continue
            key=str(durable.relative_to(ROOT));packet=json.loads(path.read_text())
            for context in contexts:
                if key not in (context.get('source_cache_paths') or context.get('baseline_files') or {}):continue
                rows,_=compose_preprint_view(packet.get('studies') or [],ROOT,True,selector={**context['selector'],'_cache_path':key},allow_additions=True)
                packet={**packet,'studies':rows}
            path.write_text(json.dumps(packet))
        manifest=renderer.build(stage,summary_dir(ROOT,source,mode),languages)
    directory=summary_dir(ROOT,source,mode)
    namespace=next(parent for parent in (directory,*directory.parents) if parent.name=='preprints_included')
    write_preprint_summary_seal(namespace,ROOT,{relative:hashlib.sha256((ROOT/relative).read_bytes()).hexdigest() for relative in ('app.py', 'source_worker.py', 'summary_outputs.py', 'scripts/build_preprint_summaries.py', 'scripts/build_cohort_summaries.py', 'runtime_code/manifest.json', 'runtime_code/MetaVir/analysis_settings_backend.py', 'runtime_code/metaPee/analysis_settings_backend.py', 'prostate_scopes.py', 'cohort_identity.py')})
    return manifest


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', choices=('metaPee', 'MetaVir'), required=True)
    parser.add_argument('--mode', choices=('yes', 'no'), required=True)
    parser.add_argument('--languages', nargs='+')
    args = parser.parse_args()
    manifest = build(args.source, args.mode, tuple(args.languages) if args.languages else None)
    print(json.dumps({'source': args.source, 'mode': args.mode, 'scopes': len(manifest['scopes']), 'keep_latest_cohort': True}))
